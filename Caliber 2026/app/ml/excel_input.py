import os
import sys
import time
import json
import shutil
from datetime import datetime, timezone
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.utils import get_column_letter

from app.ml.evaluator import evaluate

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
CONFIG_PATH = os.path.join(PROJECT_ROOT, "thresholds.json")
EXCEL_PATH = os.path.join(PROJECT_ROOT, "live_sensor_input.xlsx")
TEMP_DIR = os.path.join(PROJECT_ROOT, "scratch")

_last_valid_state: Optional[Dict[str, Any]] = None
_last_mtime: float = 0.0

def load_thresholds() -> Dict[str, Any]:
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        return json.load(f)

def read_and_evaluate_excel(file_path: Optional[str] = None) -> Dict[str, Any]:
    """
    Membaca berkas Excel live_sensor_input.xlsx, memvalidasi seluruh sel dan aturan,
    dan mengevaluasi data menggunakan fungsi murni evaluate().
    Mendukung hot-reload, file lock retry, dan state preservation.
    """
    global _last_valid_state, _last_mtime

    target_path = file_path or EXCEL_PATH
    file_name = os.path.basename(target_path)

    # 1. Pengecekan nama berkas sementara Excel (~$*)
    if file_name.startswith("~$"):
        if _last_valid_state:
            return _last_valid_state
        return {
            "file_status": {
                "file_name": file_name,
                "is_valid": False,
                "errors": [f"Berkas '{file_name}' adalah berkas lock sementara Excel dan diabaikan."]
            },
            "simulation": None
        }

    # 2. Pengecekan keberadaan berkas
    if not os.path.exists(target_path):
        return {
            "file_status": {
                "file_name": file_name,
                "is_valid": False,
                "errors": [f"Berkas '{file_name}' tidak ditemukan di lokasi yang diharapkan."]
            },
            "simulation": None
        }

    mtime = os.path.getmtime(target_path)
    mtime_iso = datetime.fromtimestamp(mtime, tz=timezone.utc).isoformat()
    read_time_iso = datetime.now(timezone.utc).isoformat()

    # 3. Salin berkas ke lokasi sementara dengan mekanisme Retry saat PermissionError
    os.makedirs(TEMP_DIR, exist_ok=True)
    temp_copy_path = os.path.join(TEMP_DIR, f"temp_read_{int(time.time()*1000)}.xlsx")

    copy_success = False
    copy_err_msg = ""
    for attempt in range(3):
        try:
            shutil.copy2(target_path, temp_copy_path)
            copy_success = True
            break
        except PermissionError as pe:
            copy_err_msg = f"Berkas sedang dikunci oleh Excel (PermissionError). Percobaan ke-{attempt+1} gagal."
            time.sleep(0.15 * (2 ** attempt))
        except Exception as e:
            copy_err_msg = f"Gagal menyalin berkas sementara: {str(e)}"
            break

    if not copy_success:
        if _last_valid_state:
            res = dict(_last_valid_state)
            res["file_status"]["is_valid"] = False
            res["file_status"]["errors"] = [copy_err_msg]
            res["file_status"]["last_read_time"] = read_time_iso
            return res
        return {
            "file_status": {
                "file_name": file_name,
                "is_valid": False,
                "errors": [copy_err_msg],
                "last_read_time": read_time_iso,
                "file_mtime": mtime_iso
            },
            "simulation": None
        }

    errors: List[str] = []
    wb = None
    try:
        # 4. Buka workbook dengan data_only=True
        try:
            wb = openpyxl.load_workbook(temp_copy_path, data_only=True)
        except Exception as e:
            errors.append(f"Berkas Excel rusak atau tidak dapat dibaca: {str(e)}")
            return _format_error_response(file_name, read_time_iso, mtime_iso, errors)

        thresholds = load_thresholds()

        # 5. Validasi Sheet PARAMETER
        if "PARAMETER" not in wb.sheetnames:
            errors.append("Sheet 'PARAMETER' tidak ditemukan di dalam berkas Excel.")
            return _format_error_response(file_name, read_time_iso, mtime_iso, errors)

        ws_param = wb["PARAMETER"]
        tag_raw = ws_param["B4"].value
        if not tag_raw:
            errors.append("Sheet PARAMETER, Baris 4, Kolom B: Tag peralatan kosong.")
            tag = "KO-3201"
        else:
            tag = str(tag_raw).strip()
            if tag not in thresholds:
                valid_tags_str = ", ".join(thresholds.keys())
                errors.append(f"Sheet PARAMETER, Baris 4, Kolom B: Tag peralatan '{tag}' tidak dikenal. Pilihan valid: {valid_tags_str}.")

        source_mode = str(ws_param["B5"].value or "GENERATOR").strip().upper()
        mode_pola = str(ws_param["B6"].value or "ramp").strip()
        scenario_desc = str(ws_param["B16"].value or "").strip()

        # 6. Pemilihan Sheet Data
        if source_mode == "INPUT_MANUAL":
            data_sheet_name = "INPUT"
            start_row = 5
        else:
            data_sheet_name = "GENERATOR"
            start_row = 4

        if data_sheet_name not in wb.sheetnames:
            errors.append(f"Sheet '{data_sheet_name}' tidak ditemukan di dalam berkas Excel.")
            return _format_error_response(file_name, read_time_iso, mtime_iso, errors)

        ws_data = wb[data_sheet_name]

        # 7. Pembacaan dan Validasi Baris Telemetri
        timestamps: List[str] = []
        load_vals: List[float] = []
        param_vals: List[float] = []

        parsed_datetimes = []

        for h in range(1, 73):
            curr_row = start_row + (h - 1)
            
            val_ts = ws_data.cell(row=curr_row, column=2).value
            val_load = ws_data.cell(row=curr_row, column=3).value
            val_param = ws_data.cell(row=curr_row, column=4).value

            # Cek sel rumus kosong (None)
            if val_ts is None:
                errors.append(f"Sheet {data_sheet_name}, Baris {curr_row}, Kolom B: Sel rumus waktu kosong (None). Buka dan simpan file di Excel agar rumus terhitung.")
            if val_load is None:
                errors.append(f"Sheet {data_sheet_name}, Baris {curr_row}, Kolom C: Sel rumus laju beban kosong (None). Buka dan simpan file di Excel agar rumus terhitung.")
            if val_param is None:
                errors.append(f"Sheet {data_sheet_name}, Baris {curr_row}, Kolom D: Sel rumus nilai parameter kosong (None). Buka dan simpan file di Excel agar rumus terhitung.")

            if errors:
                # Jika ada sel kosong, hentikan iterasi untuk menghindari cascading errors
                break

            # Validasi Waktu
            ts_dt = None
            if isinstance(val_ts, datetime):
                ts_dt = val_ts
                ts_str = val_ts.strftime("%Y-%m-%d %H:%M:%S")
            else:
                ts_str = str(val_ts).strip()
                try:
                    ts_dt = datetime.strptime(ts_str, "%Y-%m-%d %H:%M:%S")
                except ValueError:
                    try:
                        ts_dt = datetime.fromisoformat(ts_str.replace("Z", "+00:00"))
                    except Exception:
                        errors.append(f"Sheet {data_sheet_name}, Baris {curr_row}, Kolom B: Format waktu '{ts_str}' tidak valid (gunakan YYYY-MM-DD HH:MM:SS).")

            if ts_dt:
                if parsed_datetimes:
                    prev_dt = parsed_datetimes[-1]
                    # Monotonicity check
                    if ts_dt <= prev_dt:
                        errors.append(f"Sheet {data_sheet_name}, Baris {curr_row}, Kolom B: Urutan waktu tidak berurutan ({ts_str} <= {timestamps[-1]}).")
                    # Uniform interval check (3600 seconds)
                    diff_sec = int((ts_dt - prev_dt).total_seconds())
                    if diff_sec != 3600:
                        errors.append(f"Sheet {data_sheet_name}, Baris {curr_row}, Kolom B: Selang waktu tidak seragam ({diff_sec} detik, diharapkan 3600 detik).")
                parsed_datetimes.append(ts_dt)
                timestamps.append(ts_str)

            # Validasi Laju Beban
            try:
                flt_load = float(val_load)
                if flt_load < 0.0 or flt_load > 10000.0:
                    errors.append(f"Sheet {data_sheet_name}, Baris {curr_row}, Kolom C: Laju beban ({flt_load}) di luar rentang wajar (0 – 10.000).")
                load_vals.append(flt_load)
            except (ValueError, TypeError):
                errors.append(f"Sheet {data_sheet_name}, Baris {curr_row}, Kolom C: Nilai laju beban '{val_load}' bukan angka numerik.")

            # Validasi Nilai Parameter
            try:
                flt_param = float(val_param)
                if flt_param < -50.0 or flt_param > 10000.0:
                    errors.append(f"Sheet {data_sheet_name}, Baris {curr_row}, Kolom D: Nilai parameter ({flt_param}) di luar rentang wajar (-50 – 10.000).")
                param_vals.append(flt_param)
            except (ValueError, TypeError):
                errors.append(f"Sheet {data_sheet_name}, Baris {curr_row}, Kolom D: Nilai parameter '{val_param}' bukan angka numerik.")

        # 8. Jika ada error validasi, pertahankan state valid terakhir
        if errors:
            return _format_error_response(file_name, read_time_iso, mtime_iso, errors, tag, source_mode, mode_pola)

        # 9. Jalankan fungsi evaluasi murni evaluate()
        cfg = thresholds[tag]
        eval_result = evaluate(
            riwayat=param_vals,
            beban=load_vals,
            aset=tag,
            config=cfg,
            timestamps=timestamps,
            is_simulated=True
        )

        # 10. Susun Respons Berhasil
        timeseries_data = [
            {
                "hour": i + 1,
                "timestamp": timestamps[i],
                "load_rate": round(load_vals[i], 2),
                "parameter_value": round(param_vals[i], 3)
            }
            for i in range(len(param_vals))
        ]

        success_response = {
            "file_status": {
                "file_name": file_name,
                "file_path": target_path,
                "last_read_time": read_time_iso,
                "file_mtime": mtime_iso,
                "is_valid": True,
                "errors": [],
                "selected_tag": tag,
                "selected_source": source_mode,
                "selected_mode": mode_pola,
                "scenario_description": scenario_desc,
                "total_rows_read": len(param_vals)
            },
            "simulation": eval_result,
            "timeseries": timeseries_data
        }

        _last_valid_state = success_response
        _last_mtime = mtime
        return success_response

    finally:
        if wb:
            try:
                wb.close()
            except Exception:
                pass
        if os.path.exists(temp_copy_path):
            try:
                os.remove(temp_copy_path)
            except Exception:
                pass

def _format_error_response(file_name, read_time, mtime, errors, tag="KO-3201", src="GENERATOR", mode="ramp"):
    global _last_valid_state
    if _last_valid_state:
        res = json.loads(json.dumps(_last_valid_state))
        res["file_status"]["is_valid"] = False
        res["file_status"]["errors"] = errors
        res["file_status"]["last_read_time"] = read_time
        return res
    return {
        "file_status": {
            "file_name": file_name,
            "last_read_time": read_time,
            "file_mtime": mtime,
            "is_valid": False,
            "errors": errors,
            "selected_tag": tag,
            "selected_source": src,
            "selected_mode": mode,
            "scenario_description": "-"
        },
        "simulation": None,
        "timeseries": []
    }
