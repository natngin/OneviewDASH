import io
import openpyxl
import pytest
from fastapi.testclient import TestClient
from app.main import app
from app.routers.sandbox import session

client = TestClient(app)

@pytest.fixture(autouse=True)
def reset_session():
    # Reset sandbox before each test
    session.reset()
    yield
    session.reset()

def test_sandbox_status_default():
    resp = client.get("/api/sandbox/status")
    assert resp.status_code == 200
    data = resp.json()
    assert data["active"] is False
    assert data["scenario_id"] is None
    assert data["n_assets"] == 0
    assert data["n_readings"] == 0
    assert data["has_truth"] is False

def test_sandbox_template_download():
    resp = client.get("/api/sandbox/template")
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers["content-type"]
    assert "attachment; filename=\"template_uji.xlsx\"" in resp.headers["content-disposition"]
    
    # Parse downloaded excel workbook
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert set(wb.sheetnames) == {"assets", "thresholds", "readings"}
    
    ws_rd = wb["readings"]
    headers = [cell.value for cell in ws_rd[1]]
    assert headers == ["timestamp", "tag", "parameter", "value", "load_rate"]
    assert ws_rd.max_row >= 2

def test_sandbox_upload_and_results():
    # 1. Download valid template
    tpl_resp = client.get("/api/sandbox/template")
    assert tpl_resp.status_code == 200

    # 2. Upload template back to /api/sandbox/upload
    files = {"file": ("test_scenario.xlsx", tpl_resp.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    up_resp = client.post("/api/sandbox/upload", files=files)
    assert up_resp.status_code == 200
    up_data = up_resp.json()
    assert up_data["status"] == "success"
    assert up_data["n_assets"] >= 1
    assert up_data["n_readings"] >= 2
    assert "t_min" in up_data
    assert "t_max" in up_data

    # 3. Check status is now active
    stat_resp = client.get("/api/sandbox/status")
    assert stat_resp.status_code == 200
    assert stat_resp.json()["active"] is True
    assert stat_resp.json()["n_readings"] >= 2

    # 4. Run simulation
    run_resp = client.post("/api/sandbox/run", json={"as_of": "2026-10-01"})
    assert run_resp.status_code == 200
    assert run_resp.json()["status"] == "success"

    # 5. Get results
    res_resp = client.get("/api/sandbox/results?as_of=2026-10-01")
    assert res_resp.status_code == 200
    res_data = res_resp.json()
    assert "kpi" in res_data
    assert "alarms" in res_data
    assert "performance" in res_data
    assert res_data["kpi"]["totalMonitored"] >= 1
    assert len(res_data["alarms"]) >= 1

def test_sandbox_validation_errors_preserve_state():
    # Upload valid template first
    tpl_resp = client.get("/api/sandbox/template")
    files = {"file": ("valid.xlsx", tpl_resp.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    client.post("/api/sandbox/upload", files=files)
    
    orig_status = client.get("/api/sandbox/status").json()
    assert orig_status["n_readings"] >= 2

    # Now upload invalid CSV with non-numeric value
    bad_csv = "timestamp,tag,parameter,value\n2026-10-01T00:00:00Z,KO-3201,Radial Vibration,not_a_number\n"
    bad_files = {"file": ("bad.csv", bad_csv.encode("utf-8"), "text/csv")}
    bad_resp = client.post("/api/sandbox/upload", files=bad_files)
    
    assert bad_resp.status_code == 400
    detail = bad_resp.json()["detail"]
    assert "errors" in detail
    assert any("numeric" in err.lower() for err in detail["errors"])

    # Ensure previous state was NOT overwritten
    post_status = client.get("/api/sandbox/status").json()
    assert post_status["n_readings"] == orig_status["n_readings"]

def test_sandbox_ground_truth_evaluate():
    # Upload template
    tpl_resp = client.get("/api/sandbox/template")
    files = {"file": ("test.xlsx", tpl_resp.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    client.post("/api/sandbox/upload", files=files)

    # Truth CSV
    truth_csv = "scenario,tag,expected_severity,hours_warning_to_trip\nSCEN-01,KO-3201,CRITICAL,48\n"
    truth_files = {"file": ("truth.csv", truth_csv.encode("utf-8"), "text/csv")}
    eval_resp = client.post("/api/sandbox/evaluate", files=truth_files)
    assert eval_resp.status_code == 200
    eval_data = eval_resp.json()
    assert "total_scenarios" in eval_data
    assert "matched_scenarios" in eval_data
    assert "false_alarms_per_asset_day" in eval_data
    assert "scenarios" in eval_data
    
    # Requirement: NO percentage accuracy or confidence strings
    eval_text = str(eval_data).lower()
    assert "confidence" not in eval_text
    assert "% accuracy" not in eval_text

def test_sandbox_reset():
    # Upload template
    tpl_resp = client.get("/api/sandbox/template")
    files = {"file": ("test.xlsx", tpl_resp.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    client.post("/api/sandbox/upload", files=files)
    assert client.get("/api/sandbox/status").json()["active"] is True

    # Call reset
    reset_resp = client.delete("/api/sandbox/reset")
    assert reset_resp.status_code == 200
    assert reset_resp.json()["active"] is False

    # Status must be clean
    post_reset = client.get("/api/sandbox/status").json()
    assert post_reset["active"] is False
    assert post_reset["n_readings"] == 0
    assert post_reset["n_assets"] == 0
